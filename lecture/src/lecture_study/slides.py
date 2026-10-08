"""PPTX -> PDF -> bounded per-page PNGs, with native text and speaker notes."""
import subprocess
import zipfile
from pathlib import Path

import pymupdf
from pptx import Presentation

MAX_PAGES = 120


def deck_page_count(source: Path) -> int:
    if source.suffix.lower() == ".pptx":
        with zipfile.ZipFile(source) as archive:
            if len(archive.infolist()) > 10000 or sum(f.file_size for f in archive.infolist()) > 256 * 1024 * 1024:
                raise ValueError("PPTX expands beyond the 256 MiB limit")
        count = len(Presentation(source).slides)
    else:
        with pymupdf.open(source) as document:
            if document.needs_pass:
                raise ValueError("Password-protected slide decks are not supported")
            count = len(document)
    if not 1 <= count <= MAX_PAGES:
        raise ValueError(f"Slides must contain 1 to {MAX_PAGES} pages")
    return count


def render_deck(source: Path, output: Path, *, selected_pages: list[int] | None = None) -> list[dict]:
    output.mkdir(parents=True, exist_ok=True)
    notes: list[str] = []
    if source.suffix.lower() == ".pptx":
        with zipfile.ZipFile(source) as archive:
            if len(archive.infolist()) > 10000 or sum(f.file_size for f in archive.infolist()) > 256 * 1024 * 1024:
                raise ValueError("PPTX expands beyond the 256 MiB limit")
        presentation = Presentation(source)
        if not 1 <= len(presentation.slides) <= MAX_PAGES:
            raise ValueError(f"Slides must contain 1 to {MAX_PAGES} pages")
        notes = [s.notes_slide.notes_text_frame.text[:6000] if s.has_notes_slide else "" for s in presentation.slides]
        profile = (output / "office-profile").resolve().as_uri()
        subprocess.run(["libreoffice", f"-env:UserInstallation={profile}", "--headless", "--convert-to", 'pdf:impress_pdf_Export:{"ExportHiddenSlides":{"type":"boolean","value":"true"}}', "--outdir", str(output), str(source)], check=True, timeout=180, capture_output=True)
        pdf = output / f"{source.stem}.pdf"
    elif source.suffix.lower() == ".pdf":
        pdf = source
    else:
        raise ValueError("Only PPTX and PDF slide decks are supported")
    pages = []
    with pymupdf.open(pdf) as document:
        if document.needs_pass:
            raise ValueError("Password-protected slide decks are not supported")
        if not 1 <= len(document) <= MAX_PAGES:
            raise ValueError(f"Slides must contain 1 to {MAX_PAGES} pages")
        if notes and len(notes) != len(document):
            raise ValueError("PPTX conversion changed the page count; export the presentation to PDF and upload it")
        selected = selected_pages if selected_pages is not None else list(range(1, len(document) + 1))
        if not selected or selected != sorted(set(selected)) or any(p < 1 or p > len(document) for p in selected):
            raise ValueError("분석할 페이지가 첨부 장표의 범위를 벗어났습니다.")
        for number in selected:
            i, page = number - 1, document[number - 1]
            scale = 1280 / max(page.rect.width, page.rect.height)
            image = output / f"{i + 1}.png"
            page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False).save(image)
            pages.append({"page": i + 1, "text": page.get_text()[:16000], "speakerNotes": notes[i] if notes else "", "image": image})
    return pages
